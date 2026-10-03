from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
# cs_live_2ABgtGBSWFReNLdrij8BPuY5UD_xTsew
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

BG = RGBColor(0x0A, 0x14, 0x28)
CARD = RGBColor(0x13, 0x23, 0x3F)
CARD2 = RGBColor(0x1A, 0x30, 0x55)
CYAN = RGBColor(0x00, 0xE5, 0xCC)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
LIGHT = RGBColor(0xE8, 0xEE, 0xF7)
MUTED = RGBColor(0x9A, 0xAD, 0xC7)
ORANGE = RGBColor(0xFF, 0x9F, 0x43)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
RED = RGBColor(0xFF, 0x5C, 0x5C)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)
YELLOW = RGBColor(0xFD, 0xD0, 0x47)

def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

def add_bg_accent(slide):
    # top bar
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.08))
    s.fill.solid(); s.fill.fore_color.rgb = CYAN; s.line.fill.background()
    # side glow rect
    s2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(12.5), Inches(-0.6), Inches(1.8), Inches(1.8))
    s2.fill.solid(); s2.fill.fore_color.rgb = RGBColor(0x0E,0x2A,0x3A); s2.line.fill.background()

def textbox(slide, l, t, w, h, text, size=12, bold=False, color=LIGHT, align=PP_ALIGN.LEFT, font="Calibri"):
    tx = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tx.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color; p.font.name = font
    p.alignment = align
    return tx

def card(slide, l, t, w, h, fill=CARD):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = RGBColor(0x2A,0x40,0x63); s.line.width = Pt(1)
    return s

def card_text(slide, l, t, w, h, title, body, title_size=11, body_size=9, accent=CYAN):
    card(slide, l, t, w, h)
    textbox(slide, l+0.15, t+0.08, w-0.3, 0.35, title, size=title_size, bold=True, color=accent)
    tx = slide.shapes.add_textbox(Inches(l+0.15), Inches(t+0.45), Inches(w-0.3), Inches(h-0.6))
    tf = tx.text_frame; tf.word_wrap = True
    for i, line in enumerate(body):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text = line; p.font.size = Pt(body_size); p.font.color.rgb = LIGHT; p.font.name="Calibri"
        p.space_after = Pt(3)

def dashed_photo_box(slide, l, t, w, h, label, sub="Click to add picture / screenshot"):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.background(); s.line.color.rgb = CYAN; s.line.width = Pt(1.5); s.line.dash_style = 7
    textbox(slide, l, t+0.05, w, 0.35, "🖼  " + label, size=9, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    textbox(slide, l+0.1, t+0.4, w-0.2, 0.6, sub, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)

def header(slide, num, kicker, title, subtitle=""):
    textbox(slide, 0.5, 0.22, 1.0, 0.3, f"0{num}  ●  {kicker}", size=9, bold=True, color=CYAN)
    textbox(slide, 0.5, 0.5, 8.5, 0.6, title, size=26, bold=True, color=LIGHT)
    if subtitle:
        textbox(slide, 0.5, 1.1, 8.5, 0.4, subtitle, size=10, color=MUTED)
    textbox(slide, 11.9, 0.25, 0.9, 0.3, f"0{num} / 06", size=9, color=MUTED, align=PP_ALIGN.RIGHT)

def footer(slide):
    textbox(slide, 0.5, 7.0, 6, 0.3, "Team Code Wizards  •  CyberSentinel — AI Cyber Threat Detection", size=7.5, color=MUTED)
    textbox(slide, 10.5, 7.0, 2.3, 0.3, "SIH 2026  •  Confidential", size=7.5, color=MUTED, align=PP_ALIGN.RIGHT)

# ============ SLIDE 1 : INFORMATION ============
sl = prs.slides.add_slide(BLANK); bg(sl); add_bg_accent(sl)
header(sl, 1, "INFORMATION PAGE", "CYBERSENTINEL", "AI-Powered Behavioural Cyber Threat Detection  •  Monitor → Detect → Respond")
textbox(sl, 0.5, 1.55, 3.2, 0.3, "🛡  BASIC DETAILS OF PROJECT", size=10, bold=True, color=YELLOW)
card_text(sl, 0.5, 1.9, 5.6, 2.0,
    "What it is  •  PERN  •  Multi-tenant",
    ["• Monitors network / user / device behaviour; correlates 3–4 weak signals into one high-confidence detection (risk 0–100).",
     "• Opens auto-incidents (score ≥70), visual attack graph, remediation workflow + management reports.",
     "• Stack: PostgreSQL (Neon) + Express + React (Vite) + Node • Rule+heuristic now, ML-pluggable • No external AI API."],
    accent=YELLOW)
card_text(sl, 0.5, 4.05, 5.6, 1.35,
    "Why it matters — 1 weak signal = noise, 3–4 correlated = attack",
    ["• Signature-only tools miss zero-days & phishing → account-takeover chains.",
     "• SMEs have no SOC: need automated scoring, explainable reasoning + one-click Promote → Incident."],
    accent=CYAN)
# Right: PS meta with editable boxes
card(sl, 6.5, 1.55, 6.3, 5.0)
textbox(sl, 6.8, 1.65, 5.9, 0.35, "PROJECT STATEMENT  •  FILL BEFORE SUBMISSION  (editable boxes)", size=9, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
rows = [
    ("Problem Statement ID", "[  TYPE / PASTE PS-ID HERE  ]"),
    ("Problem Statement Title", "[  TYPE PROBLEM TITLE HERE — double-click to edit  ]"),
    ("Theme", "Cybersecurity  •  AI-Driven Behavioural Threat Detection"),
    ("PS Category", "Software  •  (Hardware N/A — pure web platform)"),
]
y = 2.15
for k, v in rows:
    textbox(sl, 6.85, y, 5.6, 0.25, k.upper(), size=8, bold=True, color=MUTED)
    b = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(y+0.28), Inches(5.6), Inches(0.42))
    is_edit = v.startswith("[")
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x0A,0x1A,0x33) if is_edit else CARD2
    b.line.color.rgb = YELLOW if is_edit else BLUE; b.line.width = Pt(1.2)
    textbox(sl, 6.95, y+0.32, 5.4, 0.35, v, size=9, bold=is_edit, color=YELLOW if is_edit else LIGHT)
    y += 0.82
textbox(sl, 6.85, y+0.05, 5.6, 0.25, "TEAM  •  CODE WIZARDS", size=8, bold=True, color=MUTED)
members = [
    "Pradyumn Dwivedi — Leader + Backend + Tester",
    "Rahul Gupta — UI/UX Designer",
    "Manthan Kesharwani — Protection & Middleware Manager",
    "Nayan Gupta — Database & Schema Manager",
]
tx = sl.shapes.add_textbox(Inches(6.85), Inches(y+0.3), Inches(5.6), Inches(1.1))
tf = tx.text_frame; tf.word_wrap = True
for i, m in enumerate(members):
    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
    p.text = "▸  " + m; p.font.size = Pt(8.5); p.font.color.rgb = LIGHT; p.font.name="Calibri"; p.space_after=Pt(2)
footer(sl)

# ============ SLIDE 2 : PROBLEM & SOLUTION ============
sl = prs.slides.add_slide(BLANK); bg(sl); add_bg_accent(sl)
header(sl, 2, "PROBLEM & SOLUTION", "Problem → Solution Comparison", "Signature tools miss chained attacks  •  we correlate weak signals into explainable cases")
# Problem card
card(sl, 0.5, 1.6, 3.9, 3.6, fill=RGBColor(0x2A,0x16,0x20))
textbox(sl, 0.7, 1.7, 3.5, 0.35, "❌  THE PROBLEM", size=11, bold=True, color=RED)
probs = ["• Signature-only AV misses zero-days, phishing + MFA-bypass chains.",
         "• Analysts drown in isolated alerts (Alert 1,2,3,4…) — no context.",
         "• SMEs: no SOC, slow triage, high MTTR, compliance risk.",
         "• Demos: 5 failed logins + OTP fail + new-country login seen as separate events."]
tx = sl.shapes.add_textbox(Inches(0.7), Inches(2.15), Inches(3.5), Inches(2.6))
tf = tx.text_frame; tf.word_wrap=True
for i, line in enumerate(probs):
    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
    p.text=line; p.font.size=Pt(9); p.font.color.rgb=LIGHT; p.font.name="Calibri"; p.space_after=Pt(5)
# VS badge
textbox(sl, 4.5, 3.0, 0.6, 0.6, "VS", size=16, bold=True, color=BG, align=PP_ALIGN.CENTER)
vs = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.55), Inches(3.05), Inches(0.5), Inches(0.5))
vs.fill.solid(); vs.fill.fore_color.rgb=YELLOW; vs.line.fill.background()
textbox(sl, 4.55, 3.08, 0.5, 0.45, "VS", size=14, bold=True, color=BG, align=PP_ALIGN.CENTER)
# Solution card
card(sl, 5.25, 1.6, 3.9, 3.6, fill=RGBColor(0x0D,0x2E,0x2A))
textbox(sl, 5.45, 1.7, 3.5, 0.35, "✅  OUR SOLUTION — CyberSentinel", size=11, bold=True, color=GREEN)
sols = ["• Push signals via POST /api/ingest → Normalize → Score 0–100 → Correlate 4 patterns (Phishing / Ransomware / Malware / Unauth-Access).",
        "• Score ≥55 = Detection, ≥70 = auto-Incident, ≥85 = CRITICAL + explainable reasoning & weights.",
        "• Analyst flow: Monitor (live) → Graph (attack path) → Detection (why 88?) → Promote → Incidents → Report."]
tx = sl.shapes.add_textbox(Inches(5.45), Inches(2.15), Inches(3.5), Inches(2.6))
tf = tx.text_frame; tf.word_wrap=True
for i, line in enumerate(sols):
    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
    p.text=line; p.font.size=Pt(9); p.font.color.rgb=LIGHT; p.font.name="Calibri"; p.space_after=Pt(5)
# Right column: photo + links
dashed_photo_box(sl, 9.5, 1.6, 3.3, 1.7, "ADD PHOTO: Dashboard / Attack illustration",
    "Suggestion: Monitor feed screenshot OR phishing-chain diagram from UI / internet")
card(sl, 9.5, 3.5, 3.3, 1.7)
textbox(sl, 9.7, 3.6, 2.9, 0.3, "🔗  WEBSITE  &  🎬  DEMO VIDEO  (paste before presenting)", size=8, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
for j, (lbl, val) in enumerate([("🌐 Website:", "[  paste live URL here  ]"), ("▶ Video:", "[  paste YouTube / Drive link + QR  ]")]):
    textbox(sl, 9.7, 3.95+j*0.6, 1.0, 0.3, lbl, size=8, bold=True, color=MUTED)
    b = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.65), Inches(3.95+j*0.6), Inches(1.95), Inches(0.32))
    b.fill.solid(); b.fill.fore_color.rgb=RGBColor(0x0A,0x1A,0x33); b.line.color.rgb=YELLOW; b.line.width=Pt(1)
    textbox(sl, 10.7, 3.96+j*0.6, 1.85, 0.3, val, size=7.5, color=YELLOW)
textbox(sl, 0.5, 5.35, 12.3, 1.4, "", size=9, color=LIGHT)
# comparison strip table-like
strip = ["Traditional:  Event → Signature match → Isolated alert (manual)",
         "CyberSentinel:  Network+User+Device → Multi-signal correlation → AI score → Incident (assisted)"]
for i, t in enumerate(strip):
    c = card(sl, 0.5+i*6.25, 5.45, 6.05, 0.55, fill=CARD2 if i==1 else CARD)
    textbox(sl, 0.7+i*6.25, 5.52, 5.7, 0.4, t, size=9, bold=True, color=MUTED if i==0 else CYAN)
footer(sl)

# ============ SLIDE 3 : TECHNICAL APPROACH ============
sl = prs.slides.add_slide(BLANK); bg(sl); add_bg_accent(sl)
header(sl, 3, "TECHNICAL APPROACH", "Process Flow + Stack + Implementation", "Push-only ingest  •  org-isolated  •  JWT + API-key auth")
# Flow chart boxes
flow = ["External App\n(push events)", "POST /api/ingest\n1–500 signals", "Normalize\nBASE_RISK map", "AI Score\n0–100 + factors", "Correlate\n4 patterns", "Detection +\nAuto-Incident", "React UI\n6 pages"]
x = 0.5
for i, f in enumerate(flow):
    col = CYAN if i in (1,5) else (GREEN if i==6 else BLUE)
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(1.5), Inches(0.85))
    s.fill.solid(); s.fill.fore_color.rgb = CARD2; s.line.color.rgb = col; s.line.width = Pt(1.8)
    textbox(sl, x, 1.68, 1.5, 0.8, f, size=8, bold=True, color=LIGHT, align=PP_ALIGN.CENTER)
    if i < len(flow)-1:
        arr = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x+1.52), Inches(1.98), Inches(0.24), Inches(0.18))
        arr.fill.solid(); arr.fill.fore_color.rgb = col; arr.line.fill.background()
    x += 1.78
textbox(sl, 0.5, 2.6, 12.3, 0.3, "Grouping: org | user | host   •   ≥2 distinct types OR weight ≥0.35  →  score = Σweight×100 (drop <55)  •  ≥70 auto-incident  •  confidence = min(95, 55+12×n)", size=8, color=MUTED, align=PP_ALIGN.CENTER)
# Stack row
card_text(sl, 0.5, 3.05, 4.0, 1.9, "Main Tech Stack", ["• Frontend: React 19 + Vite + Tailwind + axios + vis-network + recharts",
    "• Backend: Node + Express 4 + Zod + helmet + rate-limit + JWT/bcrypt", "• DB: PostgreSQL (Neon, pooled) — org_id isolation on every query"], accent=CYAN)
card_text(sl, 4.7, 3.05, 4.0, 1.9, "Actual Implementation (Twitter-clone demo)", ["• [PHOTO SLOT] Paste ingest code + Monitor screenshot from your Twitter clone here",
    "• Show: tweet/post with link → reportToSentinel() → Detection PHISHING ~85 → Incident", "• Auth: JWT (1h, cookie+Bearer) + cs_live_ API keys; middleware: auth + validate + error"], accent=YELLOW)
card(sl, 8.9, 3.05, 3.9, 1.9)
textbox(sl, 9.1, 3.12, 3.5, 0.3, "📸  REAL PHOTOS — REPLACE THESE BOXES", size=9, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
dashed_photo_box(sl, 9.1, 3.5, 1.8, 0.75, "UI: Monitor feed", "Paste screenshot")
dashed_photo_box(sl, 10.95, 3.5, 1.8, 0.75, "UI: Graph / Report", "Paste screenshot")
textbox(sl, 9.1, 4.35, 3.5, 0.5, "Tip: screenshots at 16:9, hide tokens. Captions: Fig.1 ingest, Fig.2 detection.", size=7.5, color=MUTED, align=PP_ALIGN.CENTER)
# Links strip
card(sl, 0.5, 5.1, 12.3, 0.9)
textbox(sl, 0.7, 5.15, 11.9, 0.3, "🔗  GITHUB  •  REPORT  •  DOCS  —  everything that doesn't fit on slides goes here", size=9, bold=True, color=YELLOW)
links = [("GitHub:", "[ paste repo URL ]"), ("Full Report:", "[ paste Drive/PDF link ]"), ("API Docs:", "documentation.md + README in repo")]
lx = 0.7
for k, v in links:
    textbox(sl, lx, 5.5, 1.1, 0.3, k, size=8, bold=True, color=CYAN)
    b = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(lx+1.0), Inches(5.5), Inches(2.6), Inches(0.32))
    b.fill.solid(); b.fill.fore_color.rgb=RGBColor(0x0A,0x1A,0x33); b.line.color.rgb=CYAN; b.line.width=Pt(1)
    textbox(sl, lx+1.05, 5.51, 2.5, 0.3, v, size=8, color=LIGHT)
    lx += 4.05
footer(sl)

# ============ SLIDE 4 : FEASIBILITY & VIABILITY ============
sl = prs.slides.add_slide(BLANK); bg(sl); add_bg_accent(sl)
header(sl, 4, "FEASIBILITY & VIABILITY  (proofread)", "Can we build it, run it, and scale it?  —  Yes.", "Pictorial: 4 pillars + risk bar  •  replace icons with project photos if needed")
pillars = [
    ("⚙", "TECHNICAL\nFEASIBLE", "PERN + Neon Postgres; no GPU / paid AI API. JWT+API keys, Zod, rate-limits already coded.", GREEN),
    ("🧩", "OPERATIONAL\n15-MIN PLUG-IN", "One helper reportToSentinel() + fire-and-forget POST. No agent, no UI change for end-users.", CYAN),
    ("💰", "ECONOMIC\n~₹0 STACK", "Neon free tier + Node + Vercel/Railway free tiers. Cost = hosting only; scales per-org.", YELLOW),
    ("⏱", "TIME & SCALE\nHACKATHON-READY", "MVP live: ingest→UI in days. Batch 1–500, 200 req/15 min; add 30-min DB window next.", PURPLE),
]
x = 0.5
for icon, title, body, col in pillars:
    card(sl, x, 1.65, 2.95, 3.1)
    textbox(sl, x, 1.75, 2.95, 0.5, icon, size=26, bold=True, color=col, align=PP_ALIGN.CENTER)
    textbox(sl, x+0.15, 2.3, 2.65, 0.5, title, size=10, bold=True, color=col, align=PP_ALIGN.CENTER)
    textbox(sl, x+0.2, 2.9, 2.55, 1.6, body, size=8.5, color=LIGHT, align=PP_ALIGN.CENTER)
    x += 3.15
# risk strip
card(sl, 0.5, 4.95, 8.6, 1.1)
textbox(sl, 0.7, 5.0, 8.2, 0.3, "⚠  RISKS  →  MITIGATIONS  (already handled / planned)", size=9, bold=True, color=ORANGE)
textbox(sl, 0.7, 5.3, 8.2, 0.6, "Batch-only correlation → query 30-min DB window  •  JWT expiry → cs_live_ keys  •  Noise/rate-limit → batch + send security events only  •  AuthZ unused → add authorize('OWNER')", size=8.5, color=LIGHT)
dashed_photo_box(sl, 9.35, 4.95, 3.45, 1.1, "ADD: deployment / cost / latency chart", "e.g. Neon + Vercel bill = ₹0, p95 score <50ms")
footer(sl)

# ============ SLIDE 5 : IMPACT & BENEFITS ============
sl = prs.slides.add_slide(BLANK); bg(sl); add_bg_accent(sl)
header(sl, 5, "IMPACT & BENEFITS", "From alert fatigue → correlated, explainable action  •  before vs after", "")
benefits = [
    ("🎯", "Fewer false positives", "1 weak signal ignored;\n3–4 correlated = case.\nConfidence 55–95.", CYAN),
    ("⚡", "Faster triage (MTTR ↓)", "Auto-incident ≥70 +\nGraph path + reasoning.\nOPEN→RESOLVED in UI.", GREEN),
    ("🏫", "SME-ready security", "No SOC needed.\nOrg-isolated, role-ready.\nSME / edu / fintech.", YELLOW),
    ("📊", "Audit & compliance", "Reports by type/severity,\ntop risky users,\nJSON export.", PURPLE),
]
x = 0.5
for icon, t, b, col in benefits:
    card(sl, x, 1.55, 2.0, 1.85)
    textbox(sl, x, 1.62, 2.0, 0.4, icon, size=20, color=col, align=PP_ALIGN.CENTER)
    textbox(sl, x+0.1, 2.1, 1.8, 0.35, t, size=9, bold=True, color=col, align=PP_ALIGN.CENTER)
    textbox(sl, x+0.1, 2.5, 1.8, 0.8, b, size=8, color=LIGHT, align=PP_ALIGN.CENTER)
    x += 2.15
# Chart: before vs after
card(sl, 0.5, 3.6, 6.4, 3.4)
textbox(sl, 0.7, 3.68, 6.0, 0.35, "📉  BEFORE vs AFTER — alerts, triage time, coverage (illustrative — replace with your measured numbers)", size=8.5, bold=True, color=YELLOW)
# chart placeholder
cd0 = ChartData()
cd0.categories = ["Isolated alerts/case", "Triage time (min)", "Chained-attack catch %"]
cd0.add_series("Before (signature-only)", (4.0, 45.0, 30.0))
cd0.add_series("After (CyberSentinel)", (1.0, 12.0, 88.0))
chart_box = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(4.05), Inches(6.0), Inches(2.6), cd0)
# populate after creation via chart data
chart = chart_box.chart
cd = ChartData()
cd.categories = ["Isolated alerts/case", "Triage time (min)", "Chained-attack catch %"]
cd.add_series("Before (signature-only)", (4.0, 45.0, 30.0))
cd.add_series("After (CyberSentinel)", (1.0, 12.0, 88.0))
chart.replace_data(cd)
chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM
chart.value_axis.has_major_gridlines = False
# photo / evidence column
card(sl, 7.15, 3.6, 5.65, 3.4)
textbox(sl, 7.35, 3.68, 5.25, 0.35, "🖼  EVIDENCE — paste your graphs / incident reports here", size=8.5, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
dashed_photo_box(sl, 7.35, 4.05, 2.65, 1.15, "ADD GRAPH: detections / severity pie", "From Report page (recharts)")
dashed_photo_box(sl, 10.15, 4.05, 2.5, 1.15, "ADD: incident case", "e.g. Phishing 85 → RESOLVED")
dashed_photo_box(sl, 7.35, 5.3, 5.3, 1.4, "ADD: before/after table, testimonials, or compliance checklist", "Keep pictorial: 1 big image > 10 bullets")
footer(sl)

# ============ SLIDE 6 : RESEARCH & REFERENCES ============
sl = prs.slides.add_slide(BLANK); bg(sl); add_bg_accent(sl)
header(sl, 6, "RESEARCH & REFERENCES", "UI map + every technology & resource  •  paste screenshots into the grid", "")
# UI grid 2x4-ish
pages = ["Home (hero+stats)", "Monitor (live feed)", "Graph (attack path)", "Detection (reasoning)", "Incidents (workflow)", "Report (charts+export)", "Login/Register", "API Keys/Admin"]
x0, y0 = 0.5, 1.6
for i, p in enumerate(pages):
    cx = x0 + (i%4)*1.98; cy = y0 + (i//4)*1.45
    dashed_photo_box(sl, cx, cy, 1.85, 1.3, f"UI: {p}", "Paste screenshot")
# right resources
card(sl, 8.65, 1.6, 4.15, 4.35)
textbox(sl, 8.85, 1.7, 3.75, 0.35, "📚  TECHNOLOGIES & RESOURCES USED", size=10, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
refs = ["• Frontend: React 19, Vite 6, Tailwind 4, axios, vis-network, recharts, lucide-react",
        "• Backend: Express 4, pg 8, jsonwebtoken 9, bcryptjs, zod, helmet, cors, rate-limit",
        "• DB/Deploy: Neon Postgres, Vercel (client), Railway/Render (API)",
        "• Docs in repo: README.md (rebuild guide), documentation.md (integration), Structure.md (STP)",
        "• Concepts: behavioural correlation, BASE_RISK scoring, 4 attack patterns, risk 0–100",
        "• References: OWASP Top-10 (phishing/A‑TO), MITRE ATT&CK (beaconing, priv‑esc, shadow‑copy)"]
tx = sl.shapes.add_textbox(Inches(8.85), Inches(2.1), Inches(3.75), Inches(3.4))
tf = tx.text_frame; tf.word_wrap=True
for i, r in enumerate(refs):
    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
    p.text=r; p.font.size=Pt(8); p.font.color.rgb=LIGHT; p.font.name="Calibri"; p.space_after=Pt(4)
card(sl, 0.5, 6.1, 12.3, 0.85)
textbox(sl, 0.7, 6.15, 11.9, 0.3, "🙏  THANK YOU  •  Team Code Wizards  —  Pradyumn (Lead/Backend/Test)  •  Rahul (UI/UX)  •  Manthan (Protection/Middleware)  •  Nayan (DB/Schema)", size=9, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
textbox(sl, 0.7, 6.5, 11.9, 0.35, "Live demo: Monitor → Graph → Detection → Promote → Incidents → Report   •   Questions welcome", size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
footer(sl)

prs.save(r"D:\Cyber\CyberSentinel_CodeWizards_6Slides.pptx")
print("saved")
